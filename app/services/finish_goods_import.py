import hashlib
import json
import unicodedata
import zipfile
from collections import Counter
from dataclasses import dataclass
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import func, select
from werkzeug.utils import secure_filename

from app.extensions import db
from app.models import (
    AuditLog,
    FinishGoodsImportBatch,
    FinishGoodsImportRow,
    FinishGoodsProfile,
    Product,
    utcnow,
)
from app.services.material_import import _validate_xlsx_container

REQUIRED_HEADERS = ("FINISH GOODS_CODE", "CATEGORY_NO", "NAME")
ALLOWED_CATEGORY = "F/G"


class FinishGoodsImportError(ValueError):
    """Raised when a Finish Goods Master import cannot be validated or applied."""


@dataclass
class ParsedFinishGood:
    row_number: int
    code: str | None
    category: str | None
    name: str | None
    result: str = "REJECTED"
    reason_code: str | None = None
    reason_detail: str | None = None


def _normalize(value, *, uppercase=False):
    if value is None:
        return None
    normalized = unicodedata.normalize("NFC", str(value)).strip()
    if not normalized:
        return None
    if any(unicodedata.category(character).startswith("C") for character in normalized):
        raise FinishGoodsImportError("Workbook values must not contain control characters.")
    return normalized.upper() if uppercase else normalized


def _safe_filename(filename):
    basename = (filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    return secure_filename(basename)[:255] or "finish-goods-master.xlsx"


def parse_finish_goods_workbook(
    file_bytes, *, maximum_bytes, maximum_rows, maximum_uncompressed_bytes
):
    if not file_bytes:
        raise FinishGoodsImportError("The uploaded workbook is empty.")
    if len(file_bytes) > maximum_bytes:
        raise FinishGoodsImportError("The workbook exceeds the permitted upload size.")
    try:
        _validate_xlsx_container(file_bytes, maximum_uncompressed_bytes)
    except ValueError as exc:
        raise FinishGoodsImportError(str(exc)) from exc
    try:
        workbook = load_workbook(
            BytesIO(file_bytes), read_only=True, data_only=False, keep_links=False
        )
    except (
        InvalidFileException,
        KeyError,
        OSError,
        TypeError,
        ValueError,
        zipfile.BadZipFile,
    ) as exc:
        raise FinishGoodsImportError(
            "The uploaded file could not be read as an .xlsx workbook."
        ) from exc
    try:
        if workbook.sheetnames != ["Sheet1"]:
            raise FinishGoodsImportError("The workbook must contain only the Sheet1 worksheet.")
        rows = list(workbook["Sheet1"].iter_rows())
        if not rows:
            raise FinishGoodsImportError("The Sheet1 worksheet is empty.")
        header = list(rows[0])
        while header and _normalize(header[-1].value) is None:
            header.pop()
        if any(cell.data_type == "f" for cell in header):
            raise FinishGoodsImportError("Header cells must not contain formulas.")
        if tuple(_normalize(cell.value, uppercase=True) for cell in header) != REQUIRED_HEADERS:
            raise FinishGoodsImportError(
                "Sheet1 headers must be exactly: FINISH GOODS_CODE, CATEGORY_NO, NAME."
            )
        data_rows = [list(row) for row in rows[1:]]
        while data_rows and all(
            cell.value is None or not str(cell.value).strip() for cell in data_rows[-1]
        ):
            data_rows.pop()
        if not data_rows:
            raise FinishGoodsImportError("The workbook contains no Finish Goods rows.")
        if len(data_rows) > maximum_rows:
            raise FinishGoodsImportError("The workbook contains more rows than permitted.")
        parsed = []
        for row_number, cells in enumerate(data_rows, start=2):
            extra = cells[3:]
            selected = cells[:3]
            while len(selected) < 3:
                selected.append(None)
            values = []
            control_error = None
            for cell, uppercase in zip(selected, (True, True, False), strict=True):
                try:
                    values.append(
                        _normalize(cell.value if cell is not None else None, uppercase=uppercase)
                    )
                except FinishGoodsImportError as exc:
                    values.append(None)
                    control_error = str(exc)
            code, category, name = values
            row = ParsedFinishGood(row_number, code, category, name)
            if any(cell.value is not None and str(cell.value).strip() for cell in extra):
                row.reason_code, row.reason_detail = (
                    "EXTRA_VALUE",
                    "Values outside the approved three columns are not accepted.",
                )
            elif control_error:
                row.reason_code, row.reason_detail = "CONTROL_CHARACTER_NOT_ALLOWED", control_error
            elif any(cell is not None and cell.data_type == "f" for cell in selected):
                row.reason_code, row.reason_detail = (
                    "FORMULA_NOT_ALLOWED",
                    "Required cells must contain values, not formulas.",
                )
            elif code is None or category is None or name is None:
                row.reason_code, row.reason_detail = (
                    "REQUIRED_VALUE_MISSING",
                    "FINISH GOODS_CODE, CATEGORY_NO, and NAME are required.",
                )
            elif len(code) > 50 or len(category) > 30 or len(name) > 200:
                row.reason_code, row.reason_detail = (
                    "VALUE_TOO_LONG",
                    "One or more values exceed the permitted field length.",
                )
            elif category != ALLOWED_CATEGORY:
                row.reason_code, row.reason_detail = (
                    "CATEGORY_NOT_ALLOWED",
                    "CATEGORY_NO must be F/G.",
                )
            parsed.append(row)
        occurrences = {}
        for row in parsed:
            if row.code:
                occurrences.setdefault(row.code, []).append(row.row_number)
        for row in parsed:
            duplicate_rows = occurrences.get(row.code, [])
            if len(duplicate_rows) > 1:
                row.reason_code = "DUPLICATE_FINISH_GOODS_CODE"
                row.reason_detail = (
                    "FINISH GOODS_CODE occurs at Excel rows "
                    + ", ".join(map(str, duplicate_rows))
                    + "."
                )
        return parsed
    finally:
        workbook.close()


def _set_counts(batch):
    counts = Counter(row.result for row in batch.rows)
    batch.total_rows = len(batch.rows)
    batch.inserted_count = counts["INSERT"]
    batch.updated_count = counts["UPDATE"]
    batch.unchanged_count = counts["UNCHANGED"]
    batch.rejected_count = counts["REJECTED"]


def _audit(event_type, batch, user_id, station_id):
    detail = json.dumps(
        {
            "batch_id": batch.id,
            "filename": batch.original_filename,
            "file_sha256": batch.file_sha256,
            "status": batch.status,
            "counts": {
                "total": batch.total_rows,
                "insert": batch.inserted_count,
                "update": batch.updated_count,
                "unchanged": batch.unchanged_count,
                "rejected": batch.rejected_count,
            },
        },
        ensure_ascii=True,
        separators=(",", ":"),
    )
    audit = AuditLog(
        event_type=event_type,
        entity_type="FINISHED_GOODS_IMPORT_BATCH",
        entity_id=str(batch.id),
        user_id=user_id,
        station_id=station_id,
        occurred_at_utc=utcnow(),
        detail=detail,
    )
    if db.session.get_bind().dialect.name == "sqlite":
        audit.id = db.session.scalar(select(func.coalesce(func.max(AuditLog.id), 0) + 1))
    db.session.add(audit)


def create_finish_goods_import_preview(
    *,
    file_bytes,
    filename,
    idempotency_key,
    user_id,
    station_id,
    maximum_bytes,
    maximum_rows,
    maximum_uncompressed_bytes,
):
    existing = db.session.scalar(
        select(FinishGoodsImportBatch).where(
            FinishGoodsImportBatch.idempotency_key == idempotency_key
        )
    )
    if existing is not None:
        return existing
    batch = FinishGoodsImportBatch(
        original_filename=_safe_filename(filename),
        file_sha256=hashlib.sha256(file_bytes).hexdigest(),
        status="PREVIEWED",
        total_rows=0,
        uploaded_by_user_id=user_id,
        uploaded_at_utc=utcnow(),
        idempotency_key=idempotency_key,
    )
    db.session.add(batch)
    try:
        parsed = parse_finish_goods_workbook(
            file_bytes,
            maximum_bytes=maximum_bytes,
            maximum_rows=maximum_rows,
            maximum_uncompressed_bytes=maximum_uncompressed_bytes,
        )
    except FinishGoodsImportError as exc:
        batch.status, batch.error_summary = "FAILED", str(exc)
        db.session.flush()
        _audit("FINISHED_GOODS_IMPORT_FAILED", batch, user_id, station_id)
        db.session.commit()
        return batch
    valid_codes = [row.code for row in parsed if row.reason_code is None]
    existing_products = {
        product.code: product
        for product in db.session.scalars(select(Product).where(Product.code.in_(valid_codes)))
    }
    for source in parsed:
        if source.reason_code is None:
            product = existing_products.get(source.code)
            if product is None:
                source.result = "INSERT"
            elif (
                product.name == source.name
                and product.finish_goods_profile is not None
                and product.finish_goods_profile.source_category_no == source.category
                and product.is_active
            ):
                source.result = "UNCHANGED"
            else:
                source.result = "UPDATE"
        batch.rows.append(
            FinishGoodsImportRow(
                row_number=source.row_number,
                code_normalized=source.code,
                category_no_normalized=source.category,
                name_normalized=source.name,
                result=source.result,
                reason_code=source.reason_code,
                reason_detail=source.reason_detail,
            )
        )
    _set_counts(batch)
    db.session.flush()
    _audit("FINISHED_GOODS_IMPORT_PREVIEWED", batch, user_id, station_id)
    db.session.commit()
    return batch


def _batch_statement(batch_id):
    return (
        select(FinishGoodsImportBatch)
        .where(FinishGoodsImportBatch.id == batch_id)
        .with_for_update()
        .with_hint(FinishGoodsImportBatch, "WITH (UPDLOCK, HOLDLOCK)", dialect_name="mssql")
    )


def apply_finish_goods_import(*, batch_id, user_id, station_id):
    batch = db.session.scalar(_batch_statement(batch_id))
    if batch is None:
        raise FinishGoodsImportError("Finish Goods import preview was not found.")
    if batch.uploaded_by_user_id != user_id:
        raise FinishGoodsImportError("Only the preview uploader can apply this import.")
    if batch.status == "APPLIED":
        return batch
    if batch.status != "PREVIEWED":
        raise FinishGoodsImportError("Only a valid preview can be applied.")
    rows = sorted(batch.rows, key=lambda row: row.row_number)
    counts = Counter(row.result for row in rows)
    if len(rows) != batch.total_rows or counts["REJECTED"] or batch.rejected_count:
        raise FinishGoodsImportError(
            "The persisted preview failed integrity validation or contains rejected rows."
        )
    codes = []
    for row in rows:
        code, category, name = (
            _normalize(row.code_normalized, uppercase=True),
            _normalize(row.category_no_normalized, uppercase=True),
            _normalize(row.name_normalized),
        )
        if (
            code != row.code_normalized
            or category != ALLOWED_CATEGORY
            or name != row.name_normalized
            or row.reason_code
            or row.result not in {"INSERT", "UPDATE", "UNCHANGED"}
        ):
            raise FinishGoodsImportError(
                "The persisted preview contains an invalid Finish Goods row."
            )
        codes.append(code)
    if len(codes) != len(set(codes)):
        raise FinishGoodsImportError("The persisted preview contains duplicate Finish Goods codes.")
    products = {
        product.code: product
        for product in db.session.scalars(select(Product).where(Product.code.in_(codes)))
    }
    changed_at = utcnow()
    for row in rows:
        product = products.get(row.code_normalized)
        if product is None:
            product = Product(
                code=row.code_normalized,
                name=row.name_normalized,
                is_active=True,
            )
            product.finish_goods_profile = FinishGoodsProfile(
                source_category_no=row.category_no_normalized,
                updated_at_utc=changed_at,
                updated_by_user_id=user_id,
            )
            db.session.add(product)
            products[product.code] = product
            row.result = "INSERT"
        elif (
            product.name != row.name_normalized
            or product.finish_goods_profile is None
            or product.finish_goods_profile.source_category_no != row.category_no_normalized
            or not product.is_active
        ):
            product.name, product.is_active = row.name_normalized, True
            if product.finish_goods_profile is None:
                product.finish_goods_profile = FinishGoodsProfile(
                    source_category_no=row.category_no_normalized
                )
            product.finish_goods_profile.source_category_no = row.category_no_normalized
            product.finish_goods_profile.updated_at_utc = changed_at
            product.finish_goods_profile.updated_by_user_id = user_id
            row.result = "UPDATE"
        else:
            row.result = "UNCHANGED"
    _set_counts(batch)
    batch.status, batch.applied_by_user_id, batch.applied_at_utc = "APPLIED", user_id, changed_at
    _audit("FINISHED_GOODS_IMPORT_APPLIED", batch, user_id, station_id)
    db.session.commit()
    return batch
